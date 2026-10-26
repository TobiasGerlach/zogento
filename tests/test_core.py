import json
import zipfile
from pathlib import Path

import pytest

from bibliothek import bibtex, extract, index, ingest, importers, pagenums
from bibliothek.store import STATUS_AUTO, STATUS_BESTAETIGT, STATUS_KEINE, Item, Library, Page


# --- BibTeX ---------------------------------------------------------------------------------

BIB = r"""% Kopf der Datei

% ------------------------------------------------
% Abschnitt Hermeneutik — Prüfvermerk
% ------------------------------------------------

@book{feestuart2015effektives,
  author    = {Fee, Gordon D. and Stuart, Douglas},
  title     = {Effektives Bibelstudium},
  year      = {2015},
  note      = {PRÜFEN: \enquote{vierte} Auflage,
               zweite Zeile},
}

% nur für Holbert
@article{holbert1981deliverance,
  author = "Holbert, John C.",
  title  = {\enquote{Deliverance Belongs to Yahweh!}: Satire in the Book of Jonah},
  journaltitle = {JSOT}, volume = 21, pages = {59--81}, year = {1981}
}
"""


def test_bibtex_roundtrip_keeps_values_verbatim():
    entries = bibtex.parse(BIB)
    assert [e.key for e in entries] == ["feestuart2015effektives", "holbert1981deliverance"]
    assert entries[0].fields["note"].startswith(r"PRÜFEN: \enquote{vierte}")
    assert entries[1].fields["volume"] == "21"
    again = bibtex.parse(bibtex.format_entry(entries[1]))[0]
    assert again.fields == entries[1].fields


def test_citekey_schema_matches_existing_keys():
    assert bibtex.citekey_stem({"author": "Fee, Gordon D. and Stuart, Douglas", "year": "2015",
                                "title": "Effektives Bibelstudium"}) == "feestuart2015effektives"
    assert bibtex.citekey_stem({"author": "Schüle, Andreas", "year": "2008", "title": "Urgeschichte"}) \
        == "schuele2008urgeschichte"
    assert bibtex.citekey_stem({"author": "{Evangelische Kirche in Deutschland}", "year": "2023",
                                "title": "Wie hältst du’s mit der Kirche?"}).startswith("evangelischekircheindeutschland2023")
    assert bibtex.citekey_stem({"title": "The Lost World"}) == "anonymojlost"


def test_import_bibtex_keeps_section_and_own_comment(tmp_path):
    lib = Library(tmp_path / "bib")
    bib = tmp_path / "refs.bib"
    bib.write_text(BIB, encoding="utf-8")
    stats = importers.import_bibtex(lib, bib, progress=lambda _: None)
    assert stats["neu"] == 2
    holbert = lib.get("holbert1981deliverance")
    assert "Prüfvermerk" in holbert.kommentar and "nur für Holbert" in holbert.kommentar
    assert "Kopf der Datei" not in lib.get("feestuart2015effektives").kommentar
    assert importers.import_bibtex(lib, bib, progress=lambda _: None)["vorhanden"] == 2


# --- Seitenzahlen ------------------------------------------------------------------------------

def _book(offset=12, n=30):
    pages = []
    for scan in range(1, n + 1):
        body = "Lorem ipsum dolor sit amet. " * 12
        printed = scan + offset
        text = f"Kapitelüberschrift {printed}\n{body}" if scan % 2 else f"{body}\n{printed}"
        pages.append(Page(scan, text if scan > 3 else "Titelei " * 40))
    return pages


def test_detect_offset_from_running_heads_and_footers():
    guess = pagenums.detect(_book(offset=-12))
    assert guess["segmente"] == [{"ab_scan": 13, "versatz": -12}]


def test_detect_jump_where_scans_are_missing():
    pages = _book(offset=5, n=40)
    # ab Scan 21 fehlen zwei Seiten im Buch → Versatz springt von +5 auf +7
    for p in pages[20:]:
        p.text = p.text.replace(str(p.scan + 5), str(p.scan + 7))
    guess = pagenums.detect(pages)
    assert [s["versatz"] for s in guess["segmente"]] == [5, 7]
    assert guess["segmente"][1]["ab_scan"] == 21
    item = Item("x", seitenzaehlung={"status": STATUS_AUTO, "segmente": guess["segmente"]})
    assert item.page_number(Page(10, "")) == "15" and item.page_number(Page(30, "")) == "37"


def test_page_ref_marks_uncertain_numbers():
    item = Item("x", seitenzaehlung={"status": STATUS_AUTO, "segmente": [{"ab_scan": 13, "versatz": -12}]})
    assert item.page_ref(Page(57, "")) == "S. 45? (Scan 57, Seitenzahl geschätzt)"
    assert item.page_ref(Page(5, "")) == "Scan 5 (Druckseite ungeklärt)"
    item.seitenzaehlung["status"] = STATUS_BESTAETIGT
    assert item.page_ref(Page(57, "")) == "S. 45 (Scan 57)"
    item.seitenzaehlung = {"status": STATUS_KEINE}
    assert item.page_ref(Page(3, "")) == "Abschn. 3"


# --- Extraktion -----------------------------------------------------------------------------------

def _pdf(path: Path, pages: list[str]) -> Path:
    """Minimal-PDF mit Textebene (Helvetica), ohne Zusatzbibliothek."""
    objs = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET"
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream")
        content = len(objs)
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content} 0 R "
                    f"/Resources << /Font << /F1 3 0 R >> >> >>")
        kids.append(f"{len(objs)} 0 R")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = "%PDF-1.4\n", []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out.encode("latin-1")))
        out += f"{i} 0 obj\n{body}\nendobj\n"
    xref = len(out.encode("latin-1"))
    out += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n" + "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    path.write_bytes(out.encode("latin-1"))
    return path


def test_pdf_pages_and_scanned_page_detection(tmp_path):
    pdf = _pdf(tmp_path / "a.pdf", ["Erste Seite mit genug Text fuer die Erkennung", "", "Dritte Seite mit reichlich Text darauf"])
    ex = extract.extract(pdf)
    assert ex.paginiert and len(ex.pages) == 3
    assert "Erste Seite" in ex.pages[0].text
    assert ex.ocr_noetig == [2]


def test_docx_odt_epub_html_pptx_plain(tmp_path):
    import docx
    d = docx.Document()
    d.add_heading("Einleitung", level=1)
    d.add_paragraph("Gattung ist eine Lesekonvention.")
    d.save(tmp_path / "a.docx")
    ex = extract.extract(tmp_path / "a.docx")
    assert not ex.paginiert and "# Einleitung" in ex.pages[0].text

    with zipfile.ZipFile(tmp_path / "a.odt", "w") as z:
        z.writestr("content.xml", '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
                   'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"><office:body><office:text>'
                   '<text:h text:outline-level="2">Teil</text:h><text:p>Absatz über Jona.</text:p>'
                   '</office:text></office:body></office:document-content>')
    assert "## Teil" in extract.extract(tmp_path / "a.odt").pages[0].text

    with zipfile.ZipFile(tmp_path / "a.epub", "w") as z:
        z.writestr("META-INF/container.xml", '<container><rootfiles><rootfile full-path="OEBPS/c.opf"/></rootfiles></container>')
        z.writestr("OEBPS/c.opf", '<package xmlns:dc="http://purl.org/dc/elements/1.1/"><metadata><dc:title>Buch</dc:title>'
                   '</metadata><manifest><item id="k1" href="k1.xhtml"/></manifest><spine><itemref idref="k1"/></spine></package>')
        z.writestr("OEBPS/k1.xhtml", "<html><body><h1>Kapitel</h1><p>Text &amp; mehr</p><script>x()</script></body></html>")
    ex = extract.extract(tmp_path / "a.epub")
    assert ex.meta["title"] == "Buch" and "Text & mehr" in ex.pages[0].text and "x()" not in ex.pages[0].text

    with zipfile.ZipFile(tmp_path / "a.pptx", "w") as z:
        for n in (2, 1):
            z.writestr(f"ppt/slides/slide{n}.xml", f'<p:sld xmlns:p="p" xmlns:a="a"><a:p><a:r><a:t>Folie {n}</a:t></a:r></a:p></p:sld>')
    ex = extract.extract(tmp_path / "a.pptx")
    assert [p.text for p in ex.pages] == ["Folie 1", "Folie 2"]

    (tmp_path / "b.txt").write_text("Seite eins\fSeite zwei\f", encoding="utf-8")
    ex = extract.extract(tmp_path / "b.txt")
    assert ex.paginiert and [p.text for p in ex.pages] == ["Seite eins", "Seite zwei"]

    (tmp_path / "c.rtf").write_text(r"{\rtf1\ansi Hallo {\b Welt}}", encoding="utf-8")
    assert "Hallo Welt" in extract.extract(tmp_path / "c.rtf").pages[0].text


# --- Aufnahme und Suche ---------------------------------------------------------------------------

def test_add_file_index_and_search(tmp_path):
    lib = Library(tmp_path / "bib")
    pdf = _pdf(tmp_path / "jona.pdf", ["Jona flieht nach Tarschisch vor dem Auftrag", "Der Fisch verschlingt den Propheten"])
    item = ingest.add_file(lib, pdf, felder={"author": "Muster, Max", "year": "2020", "title": "Jona lesen"},
                           tags=["Jona"], run_ocr=False)
    assert item.citekey == "muster2020jona"
    assert (lib.folder(item.citekey) / "jona.pdf").exists()
    md = (lib.folder(item.citekey) / "text.md").read_text(encoding="utf-8")
    assert "[[S. 1]]" not in md and "Scan 1" in md    # PDF ohne Labels: Druckseite ungeklärt
    hits = index.search(lib, "Propheten", mode="text")
    assert hits and hits[0].citekey == item.citekey and hits[0].scan == 2
    assert index.search(lib, "Prophet", mode="text")      # Präfixsuche
    assert index.search(lib, "Propheten", mode="text", keys=set()) == []
    with pytest.raises(ingest.Duplicate):
        ingest.add_file(lib, pdf, felder={"author": "Muster, Max", "year": "2020", "title": "Jona lesen"}, run_ocr=False)


def test_missing_metadata_is_flagged_not_invented(tmp_path):
    lib = Library(tmp_path / "bib")
    (tmp_path / "notiz_zu_jona.md").write_text("# Jona\n\nText", encoding="utf-8")
    item = ingest.add_file(lib, tmp_path / "notiz_zu_jona.md", run_ocr=False)
    assert set(item.unvollstaendig) == {"author", "year"}
    assert item.felder["title"] == "notiz zu jona"
    assert any("Dateinamen" in h for h in item.text["hinweise"])


def test_ocr_output_import_attaches_to_existing_entry(tmp_path):
    lib = Library(tmp_path / "bib")
    lib.save(Item("feestuart2015effektives", felder={"title": "Effektives Bibelstudium", "year": "2015",
                                                    "author": "Fee, Gordon D. and Stuart, Douglas"}))
    out = tmp_path / "ocr"
    (out / "text").mkdir(parents=True)
    entries = []
    for i, page in enumerate(_book(offset=-2, n=12), 1):
        (out / "text" / f"{i}.txt").write_text(page.text, encoding="utf-8")
        entries.append({"source": f"IMG_{i}.HEIC", "pages": [{"text_file": f"text/{i}.txt", "confidence": 95.0}]})
    (out / "report.json").write_text(json.dumps({"title": "Effektives Bibelstudium", "complete": True,
                                                  "entries": entries}), encoding="utf-8")
    item = importers.import_ocr_dir(lib, out, progress=lambda _: None)
    assert item.citekey == "feestuart2015effektives"
    assert item.text["quelle"] == "ocr" and item.text["konfidenz"] == 95.0
    assert item.seitenzaehlung["status"] == STATUS_AUTO
    assert item.seitenzaehlung["segmente"] == [{"ab_scan": 3, "versatz": -2}]


def test_chunking_keeps_all_text():
    text = ("Satz eins ist hier. " * 80).strip()
    chunks = index.chunk_text(text, size=300, overlap=50)
    assert all(len(c) <= 360 for c in chunks)
    assert chunks[0].startswith("Satz eins") and chunks[-1].endswith("hier.")


def test_guess_title_rejects_scanner_and_word_junk():
    from bibliothek.ingest import guess_title
    assert guess_title("KM_364e-20201007124129", Path("Frische_Theologie_Herrschaft.pdf"))[0] == "Frische Theologie Herrschaft"
    assert guess_title("Microsoft Word - THE 5305 Leitfaden.docx", Path("x.pdf"))[0] == "THE 5305 Leitfaden"
    assert guess_title(None, Path("Newbigin_Griechen_Torheit-1.pdf"))[0] == "Newbigin Griechen Torheit"
    assert guess_title("Drowning Jonah", Path("x.pdf"))[0] == "Drowning Jonah"
